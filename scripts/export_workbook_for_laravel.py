#!/usr/bin/env python3
"""Validate the canonical workbook and create a deterministic offline Laravel manifest.

This module is intentionally offline-only. It does not import or call any hosted
database client. Canonical entry identifiers remain compatible with the legacy
keyset MD5, while SHA-256 digests provide modern reproducibility evidence.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
from pathlib import Path
import re
import sys
import unicodedata
from typing import Any, Iterable

from openpyxl import load_workbook

EXPECTED_SHEETS = [
    "README_PRODUCTION",
    "TEMPLATES",
    "GLOBAL",
    "TENANTS_MASTER",
    "DEVICES",
    "SERVICES",
    "APP_CONFIG_EXPORT",
]
EXPECTED_COUNTS = {
    "global": 1,
    "template": 3,
    "tenant": 100,
    "service": 140,
    "device": 108,
}
EXPECTED_TEMPLATES = {"BASIC", "PRO", "ENTERPRISE"}
EXPECTED_TOTAL = sum(EXPECTED_COUNTS.values())
# Kept only as a compatibility oracle. SHA-256 is the integrity digest.
EXPECTED_KEYSET_MD5 = "1fd3d64219ecc7c333453c20da80e295"
DEFAULT_WORKBOOK = (
    Path(__file__).resolve().parents[1]
    / "config/FINAL_Production_Excel_A_WebC_Template_100Tenants.xlsx"
)
DATA_SHEETS = {
    "TEMPLATES",
    "GLOBAL",
    "TENANTS_MASTER",
    "DEVICES",
    "SERVICES",
    "APP_CONFIG_EXPORT",
}
TRUE_VALUES = {"true", "1", "yes", "active", "enabled"}
FALSE_VALUES = {"false", "0", "no", "inactive", "disabled"}
KEY_ID_RE = re.compile(r"^[^:\s]+$", re.UNICODE)
SERVICE_KEY_RE = re.compile(r"^tenant:([^:\s]+):service:(pos|akuntansi)$")
TENANT_KEY_RE = re.compile(r"^tenant:([^:\s]+)$")
DEVICE_KEY_RE = re.compile(r"^device:([^:\s]+)$")
TEMPLATE_KEY_RE = re.compile(r"^template:(BASIC|PRO|ENTERPRISE)$")


def _nfc(value: str) -> str:
    return unicodedata.normalize("NFC", value)


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def required_text(value: Any, field: str) -> str:
    """Require a canonical text identifier; never coerce numbers/dates to strings."""
    if not isinstance(value, str):
        raise ValueError(f"{field} must be stored as text")
    if value != value.strip():
        raise ValueError(f"{field} must not contain leading/trailing whitespace")
    normalized = _nfc(value)
    if not normalized or not KEY_ID_RE.fullmatch(normalized):
        raise ValueError(f"{field} is blank or contains a prohibited separator/whitespace")
    return normalized


def required_key_text(value: Any, field: str) -> str:
    """Validate key text before applying the canonical key grammar."""
    if not isinstance(value, str):
        raise ValueError(f"{field} must be stored as text")
    if value != value.strip():
        raise ValueError(f"{field} must not contain leading/trailing whitespace")
    normalized = _nfc(value)
    if not normalized or any(character.isspace() for character in normalized):
        raise ValueError(f"{field} is blank or contains whitespace")
    return normalized


def _required_header_sets(sheet_name: str) -> tuple[set[str], ...]:
    required: dict[str, tuple[set[str], ...]] = {
        "TEMPLATES": ({"template_id", "template"},),
        "TENANTS_MASTER": ({"tenant_id"}, {"template_id"}),
        "DEVICES": ({"device_id"}, {"tenant_id"}, {"is_active", "active", "status"}),
        "SERVICES": ({"tenant_id"}, {"service"}, {"enabled"}),
        "APP_CONFIG_EXPORT": ({"key"}, {"value_jsonb"}, {"type"}, {"source_sheet"}, {"sync_status"}),
        # GLOBAL is format-checked but its business columns are not reinterpreted
        # here because DOC_00 defines the export keys, not a fixed GLOBAL row schema.
        "GLOBAL": (),
    }
    return required.get(sheet_name, ())


def table_rows(ws: Any, sheet_name: str) -> list[dict[str, Any]]:
    """Read a tabular sheet without silently dropping duplicate/blank headers."""
    values = list(ws.iter_rows(values_only=True))
    if not values:
        raise ValueError(f"{sheet_name} sheet is empty")

    raw_headers = list(values[0])
    populated = [i for i, value in enumerate(raw_headers) if not _is_blank(value)]
    if not populated:
        raise ValueError(f"{sheet_name} has no header row")
    last_header = max(populated)

    headers: list[str] = []
    seen_headers: set[str] = set()
    for index, raw in enumerate(raw_headers[: last_header + 1]):
        if _is_blank(raw):
            raise ValueError(f"{sheet_name} has a blank header at column {index + 1}")
        if not isinstance(raw, str):
            raise ValueError(f"{sheet_name} headers must be text")
        header = _nfc(raw.strip()).lower()
        if not header:
            raise ValueError(f"{sheet_name} has a blank header at column {index + 1}")
        if header in seen_headers:
            raise ValueError(f"{sheet_name} has duplicate header {header!r}")
        seen_headers.add(header)
        headers.append(header)

    for header_set in _required_header_sets(sheet_name):
        if not (seen_headers & header_set):
            raise ValueError(
                f"{sheet_name} is missing required header; expected one of {sorted(header_set)}"
            )

    rows: list[dict[str, Any]] = []
    for row_number, raw_row in enumerate(values[1:], start=2):
        raw_row = list(raw_row)
        if any(not _is_blank(value) for value in raw_row[last_header + 1 :]):
            raise ValueError(
                f"{sheet_name} row {row_number} has data beyond the final header"
            )
        cells = raw_row[: last_header + 1]
        if not any(not _is_blank(value) for value in cells):
            continue
        cells += [None] * (len(headers) - len(cells))
        rows.append(dict(zip(headers, cells)))
    return rows


def _unique_object_pairs(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for raw_key, value in pairs:
        key = _nfc(raw_key)
        if key in result:
            raise ValueError("duplicate JSON object property")
        result[key] = value
    return result


def _reject_json_constant(_: str) -> None:
    raise ValueError("non-standard JSON numeric constant is prohibited")


def normalize_json(value: Any, path: str = "value") -> Any:
    """Normalize text to NFC and recursively sort object keys for stable JSON."""
    if isinstance(value, dict):
        normalized: dict[str, Any] = {}
        for raw_key, raw_value in value.items():
            if not isinstance(raw_key, str):
                raise ValueError(f"{path} contains a non-text JSON object key")
            key = _nfc(raw_key)
            if key in normalized:
                raise ValueError(f"{path} contains colliding keys after Unicode normalization")
            normalized[key] = normalize_json(raw_value, f"{path}.{key}")
        return {key: normalized[key] for key in sorted(normalized)}
    if isinstance(value, list):
        return [normalize_json(item, f"{path}[]") for item in value]
    if isinstance(value, str):
        return _nfc(value)
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError(f"{path} contains a non-finite number")
    if value is None or isinstance(value, (bool, int, float)):
        return value
    raise ValueError(f"{path} contains an unsupported value type")


def parse_value(raw: Any, key: str) -> dict[str, Any]:
    if isinstance(raw, dict):
        parsed = raw
    elif isinstance(raw, str):
        try:
            parsed = json.loads(
                raw,
                object_pairs_hook=_unique_object_pairs,
                parse_constant=_reject_json_constant,
            )
        except (json.JSONDecodeError, ValueError) as exc:
            raise ValueError(f"{key}: value_jsonb is not valid strict JSON") from exc
    else:
        raise ValueError(f"{key}: value_jsonb must be a JSON object stored as text")

    if not isinstance(parsed, dict):
        raise ValueError(f"{key}: value_jsonb must decode to a JSON object")
    parsed = normalize_json(parsed, key)
    version = parsed.get("version", 1)
    if isinstance(version, bool) or not isinstance(version, int) or version < 1:
        raise ValueError(f"{key}: version must be a positive JSON integer")
    parsed["version"] = version
    return {property_name: parsed[property_name] for property_name in sorted(parsed)}


def classify_key(key: str) -> tuple[str, str | None, str | None]:
    """Parse a canonical key exactly; prefix-only counting is deliberately rejected."""
    if key == "global":
        return "global", None, None
    match = TEMPLATE_KEY_RE.fullmatch(key)
    if match:
        return "template", match.group(1), None
    match = SERVICE_KEY_RE.fullmatch(key)
    if match:
        return "service", match.group(1), match.group(2)
    match = TENANT_KEY_RE.fullmatch(key)
    if match:
        return "tenant", match.group(1), None
    match = DEVICE_KEY_RE.fullmatch(key)
    if match:
        return "device", match.group(1), None
    raise ValueError("APP_CONFIG_EXPORT contains a key outside the canonical key grammar")


def parse_bool(value: Any, field: str) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, int) and not isinstance(value, bool) and value in (0, 1):
        return bool(value)
    if isinstance(value, str):
        normalized = _nfc(value.strip()).lower()
        if normalized in TRUE_VALUES:
            return True
        if normalized in FALSE_VALUES:
            return False
    raise ValueError(f"{field} must be an explicit boolean / 0 / 1 / known status")


def _validate_formula_and_error_cells(wb: Any) -> None:
    """Never consume stale formula caches as canonical configuration values."""
    for sheet_name in DATA_SHEETS:
        ws = wb[sheet_name]
        for row in ws.iter_rows(values_only=False):
            for cell in row:
                if cell.data_type == "f" or (
                    isinstance(cell.value, str) and cell.value.startswith("=")
                ):
                    raise ValueError(
                        f"{sheet_name} contains a formula at {cell.coordinate}; "
                        "configuration source cells must be literal values"
                    )
                if cell.data_type == "e":
                    raise ValueError(
                        f"{sheet_name} contains an Excel error cell at {cell.coordinate}"
                    )


def _assert_expected_keyset(actual: set[str], expected: set[str]) -> None:
    if actual != expected:
        missing_count = len(expected - actual)
        extra_count = len(actual - expected)
        raise ValueError(
            "APP_CONFIG_EXPORT does not match the relational sheets "
            f"(missing={missing_count}, extra={extra_count})"
        )


def build_manifest(
    workbook: Path,
    *,
    expected_keyset_md5: str | None = EXPECTED_KEYSET_MD5,
) -> dict[str, Any]:
    workbook = Path(workbook)
    if not workbook.is_file():
        raise FileNotFoundError(f"Workbook does not exist: {workbook}")

    # data_only=False is intentional: formula cells are rejected, not replaced
    # with potentially stale/empty cached results from a spreadsheet application.
    wb = load_workbook(workbook, data_only=False, read_only=True)
    try:
        if wb.sheetnames != EXPECTED_SHEETS:
            raise ValueError(
                f"Unexpected sheet order/layout; expected the seven canonical sheets "
                f"in DOC_00 (actual_count={len(wb.sheetnames)})"
            )
        if any(wb[name].sheet_state != "visible" for name in EXPECTED_SHEETS):
            raise ValueError("All canonical workbook sheets must be visible")
        if not any(not _is_blank(value) for row in wb["README_PRODUCTION"].iter_rows(values_only=True) for value in row):
            raise ValueError("README_PRODUCTION must contain workbook instructions or metadata")

        _validate_formula_and_error_cells(wb)
        templates = table_rows(wb["TEMPLATES"], "TEMPLATES")
        global_rows = table_rows(wb["GLOBAL"], "GLOBAL")
        tenant_rows = table_rows(wb["TENANTS_MASTER"], "TENANTS_MASTER")
        device_rows = table_rows(wb["DEVICES"], "DEVICES")
        service_rows = table_rows(wb["SERVICES"], "SERVICES")
        export_rows = table_rows(wb["APP_CONFIG_EXPORT"], "APP_CONFIG_EXPORT")
        if not global_rows:
            raise ValueError("GLOBAL must contain at least one data row")

        by_key: dict[str, dict[str, Any]] = {}
        counts = {category: 0 for category in EXPECTED_COUNTS}
        for row in export_rows:
            key = required_key_text(row.get("key"), "APP_CONFIG_EXPORT.key")
            if key in by_key:
                raise ValueError("APP_CONFIG_EXPORT contains a duplicate key")
            kind, _, _ = classify_key(key)
            expected_source_sheet = {
                "global": "GLOBAL",
                "template": "TEMPLATES",
                "tenant": "TENANTS_MASTER",
                "service": "SERVICES",
                "device": "DEVICES",
            }[kind]
            row_type = required_text(
                row.get("type"), "APP_CONFIG_EXPORT.type"
            ).lower()
            if row_type != kind:
                raise ValueError(
                    "APP_CONFIG_EXPORT type does not match canonical key scope"
                )
            source_sheet = required_text(
                row.get("source_sheet"), "APP_CONFIG_EXPORT.source_sheet"
            ).upper()
            if source_sheet != expected_source_sheet:
                raise ValueError(
                    "APP_CONFIG_EXPORT source_sheet does not match canonical key scope"
                )
            counts[kind] += 1
            status = row.get("sync_status")
            if status is not None and str(status).strip().upper() not in {"", "READY", "ACTIVE"}:
                raise ValueError("APP_CONFIG_EXPORT contains a non-ready sync_status")
            by_key[key] = {"key": key, "config": parse_value(row.get("value_jsonb"), key)}

        if counts != EXPECTED_COUNTS or len(by_key) != EXPECTED_TOTAL:
            raise ValueError(
                f"Workbook entry count mismatch: expected {EXPECTED_COUNTS}/{EXPECTED_TOTAL}, "
                f"actual {counts}/{len(by_key)}"
            )

        template_ids: list[str] = []
        for row in templates:
            raw_template = row.get("template_id") if row.get("template_id") is not None else row.get("template")
            template_id = required_text(raw_template, "TEMPLATES.template_id").upper()
            if template_id not in EXPECTED_TEMPLATES:
                raise ValueError("TEMPLATES contains an unsupported template identifier")
            template_ids.append(template_id)
        if len(template_ids) != len(set(template_ids)) or set(template_ids) != EXPECTED_TEMPLATES:
            raise ValueError("TEMPLATES must contain exactly BASIC, PRO, and ENTERPRISE once each")

        tenant_templates: dict[str, str] = {}
        for row in tenant_rows:
            tenant_id = required_text(row.get("tenant_id"), "TENANTS_MASTER.tenant_id")
            raw_template = row.get("template_id")
            template_id = required_text(raw_template, f"TENANTS_MASTER.{tenant_id}.template_id").upper()
            if tenant_id in tenant_templates:
                raise ValueError("TENANTS_MASTER contains duplicate tenant identifiers")
            if template_id not in EXPECTED_TEMPLATES:
                raise ValueError("TENANTS_MASTER contains a missing/invalid template parent")
            tenant_templates[tenant_id] = template_id
        if len(tenant_templates) != EXPECTED_COUNTS["tenant"]:
            raise ValueError(
                f"TENANTS_MASTER must contain 100 unique tenants; got {len(tenant_templates)}"
            )

        tenants: list[dict[str, Any]] = []
        expected_keys = {"global"}
        expected_keys.update(f"template:{template_id}" for template_id in EXPECTED_TEMPLATES)
        expected_keys.update(f"tenant:{tenant_id}" for tenant_id in tenant_templates)
        for tenant_id, template_id in tenant_templates.items():
            key = f"tenant:{tenant_id}"
            if key not in by_key:
                raise ValueError("TENANTS_MASTER references a missing tenant config key")
            tenants.append({
                "tenant_id": tenant_id,
                "template_id": template_id,
                "config": by_key[key]["config"],
            })
        if set(template_ids) != {
            classify_key(key)[1] for key in by_key if classify_key(key)[0] == "template"
        }:
            raise ValueError("Template rows and canonical export keys do not match")

        services: list[dict[str, Any]] = []
        service_pairs: set[tuple[str, str]] = set()
        for row in service_rows:
            tenant_id = required_text(row.get("tenant_id"), "SERVICES.tenant_id")
            if tenant_id not in tenant_templates:
                raise ValueError("SERVICES references an unknown tenant parent")
            service = required_text(row.get("service"), "SERVICES.service").lower()
            if service not in {"pos", "akuntansi"}:
                raise ValueError("SERVICES contains an unsupported service identifier")
            pair = (tenant_id, service)
            if pair in service_pairs:
                raise ValueError("SERVICES contains a duplicate tenant/service pair")
            service_pairs.add(pair)
            key = f"tenant:{tenant_id}:service:{service}"
            provided_key = row.get("app_config_key")
            if not _is_blank(provided_key):
                provided_key = required_key_text(provided_key, "SERVICES.app_config_key")
                if provided_key != key:
                    raise ValueError("SERVICES.app_config_key crosses tenant/service scope")
            if key not in by_key:
                raise ValueError("SERVICES references a missing canonical service config key")
            expected_keys.add(key)
            services.append({
                "tenant_id": tenant_id,
                "service": service,
                "enabled": parse_bool(row.get("enabled"), "SERVICES.enabled"),
                "config": by_key[key]["config"],
            })
        if len(services) != EXPECTED_COUNTS["service"]:
            raise ValueError(f"SERVICES must contain 140 unique rows; got {len(services)}")
        export_service_keys = {
            key for key in by_key if classify_key(key)[0] == "service"
        }
        mapped_service_keys = {
            f"tenant:{tenant_id}:service:{service}" for tenant_id, service in service_pairs
        }
        if export_service_keys != mapped_service_keys:
            raise ValueError("Service export keys and SERVICES rows do not match exactly")

        devices: list[dict[str, Any]] = []
        seen_devices: set[str] = set()
        for row in device_rows:
            device_id = required_text(row.get("device_id"), "DEVICES.device_id")
            tenant_id = required_text(row.get("tenant_id"), "DEVICES.tenant_id")
            if device_id in seen_devices:
                raise ValueError("DEVICES contains duplicate device identifiers")
            if tenant_id not in tenant_templates:
                raise ValueError("DEVICES references an unknown tenant parent")
            key = f"device:{device_id}"
            provided_key = row.get("app_config_key")
            if not _is_blank(provided_key):
                provided_key = required_key_text(provided_key, "DEVICES.app_config_key")
                if provided_key != key:
                    raise ValueError("DEVICES.app_config_key crosses device scope")
            if key not in by_key:
                raise ValueError("DEVICES references a missing canonical device config key")
            activity = row.get("is_active")
            if activity is None:
                activity = row.get("active")
            if activity is None:
                activity = row.get("status")
            devices.append({
                "device_id": device_id,
                "tenant_id": tenant_id,
                "active": parse_bool(activity, "DEVICES.is_active"),
                "config": by_key[key]["config"],
            })
            seen_devices.add(device_id)
            expected_keys.add(key)
        if len(devices) != EXPECTED_COUNTS["device"]:
            raise ValueError(f"DEVICES must contain 108 unique rows; got {len(devices)}")
        export_device_keys = {key for key in by_key if classify_key(key)[0] == "device"}
        mapped_device_keys = {f"device:{device_id}" for device_id in seen_devices}
        if export_device_keys != mapped_device_keys:
            raise ValueError("Device export keys and DEVICES rows do not match exactly")

        _assert_expected_keyset(set(by_key), expected_keys)
        keyset_md5 = hashlib.md5("\n".join(sorted(by_key)).encode("utf-8")).hexdigest()
        if expected_keyset_md5 is not None and keyset_md5 != expected_keyset_md5:
            raise ValueError("Canonical keyset MD5 compatibility check failed")

        # Stable list order prevents Excel row order from changing the exported
        # projection. JSON objects are normalized/sorted by normalize_json().
        tenants.sort(key=lambda item: item["tenant_id"])
        services.sort(key=lambda item: (item["tenant_id"], item["service"]))
        devices.sort(key=lambda item: item["device_id"])
        template_payload = [
            {"id": template_id, "config": by_key[f"template:{template_id}"]["config"]}
            for template_id in sorted(EXPECTED_TEMPLATES)
        ]
        manifest: dict[str, Any] = {
            "schema_version": 1,
            "source_workbook": workbook.name,
            "source_workbook_sha256": hashlib.sha256(workbook.read_bytes()).hexdigest(),
            "counts": {name: counts[name] for name in EXPECTED_COUNTS},
            "keyset_md5": keyset_md5,
            "keyset_sha256": hashlib.sha256("\n".join(sorted(by_key)).encode("utf-8")).hexdigest(),
            "global_config": by_key["global"]["config"],
            "templates": template_payload,
            "tenants": tenants,
            "services": services,
            "devices": devices,
        }
        manifest["manifest_sha256"] = hashlib.sha256(canonical_json_bytes(manifest)).hexdigest()
        return manifest
    finally:
        wb.close()


def canonical_json_bytes(value: dict[str, Any]) -> bytes:
    """UTF-8, compact JSON, fixed insertion order, LF-terminated."""
    return (
        json.dumps(
            value,
            ensure_ascii=False,
            sort_keys=False,
            separators=(",", ":"),
            allow_nan=False,
        )
        + "\n"
    ).encode("utf-8")


def manifest_bytes(manifest: dict[str, Any]) -> bytes:
    """Serialize the manifest; its digest is calculated without manifest_sha256."""
    return canonical_json_bytes(manifest)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--output", type=Path, default=Path("build/central-config-import.json"))
    parser.add_argument(
        "--validate-only",
        action="store_true",
        help="Validate locally without writing an import manifest or accessing any service",
    )
    args = parser.parse_args(argv)
    try:
        manifest = build_manifest(args.workbook)
        print(
            "PASS: canonical workbook contract; "
            f"total={sum(manifest['counts'].values())}; counts={manifest['counts']}; "
            f"keyset_md5={manifest['keyset_md5']}; "
            f"keyset_sha256={manifest['keyset_sha256']}; "
            f"manifest_sha256={manifest['manifest_sha256']}; "
            f"source_workbook_sha256={manifest['source_workbook_sha256']}"
        )
        if args.validate_only:
            return 0
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_bytes(manifest_bytes(manifest))
        print(f"Wrote deterministic local manifest: {args.output}")
        return 0
    except Exception as exc:
        print(f"ERROR: workbook validation/export failed: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
