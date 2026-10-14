#!/usr/bin/env python3
"""Validate and sync the FINAL Pos workbook to Supabase through PostgREST."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import sys
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

import requests
from openpyxl import load_workbook

EXCEL_PATH = Path(
    os.environ.get(
        "FINAL_POS_EXCEL",
        "config/FINAL_Production_Excel_A_WebC_Template_100Tenants.xlsx",
    )
)
EXPECTED_SHEETS = [
    "README_PRODUCTION",
    "TEMPLATES",
    "GLOBAL",
    "TENANTS_MASTER",
    "DEVICES",
    "SERVICES",
    "APP_CONFIG_EXPORT",
]
EXPECTED_COUNTS = {
    "global": 1,
    "template": 3,
    "tenant": 100,
    "service": 140,
    "device": 108,
}
EXPECTED_TEMPLATES = {"BASIC", "PRO", "ENTERPRISE"}
EXPECTED_KEYS = 352
EXPECTED_KEYSET_MD5 = "1fd3d64219ecc7c333453c20da80e295"
BATCH_SIZE = 100
REQUEST_TIMEOUT_SECONDS = 30


def _keyset_md5(keys: Iterable[str]) -> str:
    """Return the canonical order-independent MD5 fingerprint of newline-joined sorted keys.

    This is an inventory identity check, not a cryptographic integrity/signature mechanism.
    """
    normalized = sorted(str(key) for key in keys)
    return hashlib.md5("\n".join(normalized).encode("utf-8")).hexdigest()


def _assert_canonical_key_manifest(label: str, keys: Iterable[str]) -> str:
    normalized = [str(key) for key in keys]
    if len(normalized) != EXPECTED_KEYS:
        raise ValueError(
            f"{label} must contain exactly {EXPECTED_KEYS} canonical keys; got {len(normalized)}"
        )
    if len(set(normalized)) != EXPECTED_KEYS:
        raise ValueError(f"{label} contains duplicate keys")
    digest = _keyset_md5(normalized)
    if digest != EXPECTED_KEYSET_MD5:
        raise ValueError(
            f"{label} key manifest MD5 mismatch: expected {EXPECTED_KEYSET_MD5}, got {digest}"
        )
    return digest


@dataclass(frozen=True)
class WorkbookPayloads:
    app_config: list[dict[str, Any]]
    tenants: list[dict[str, Any]]
    devices: list[dict[str, Any]]
    services: list[dict[str, Any]]
    counts: dict[str, int]


def _rows_as_dicts(ws) -> list[dict[str, Any]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    result: list[dict[str, Any]] = []
    for row in rows[1:]:
        if all(value is None for value in row):
            continue
        result.append({headers[index]: row[index] for index in range(len(headers))})
    return result


def _normalized_row(row: dict[str, Any]) -> dict[str, Any]:
    return {
        str(key).strip().lower().replace(" ", "_"): value
        for key, value in row.items()
        if str(key).strip()
    }


def _pick(row: dict[str, Any], *names: str) -> Any:
    normalized = _normalized_row(row)
    for name in names:
        key = name.strip().lower().replace(" ", "_")
        value = normalized.get(key)
        if value is not None and str(value).strip() != "":
            return value
    return None


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _required_text(row: dict[str, Any], label: str, *aliases: str) -> str:
    value = _text(_pick(row, *aliases))
    if not value:
        raise ValueError(f"{label} is required; accepted columns: {aliases}")
    return value


def _bool(value: Any, *, default: bool) -> bool:
    if value is None or _text(value) == "":
        return default
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    normalized = _text(value).lower()
    if normalized in {"true", "1", "yes", "y", "aktif", "active", "enabled"}:
        return True
    if normalized in {"false", "0", "no", "n", "nonaktif", "inactive", "disabled"}:
        return False
    raise ValueError(f"Cannot parse boolean value: {value!r}")


def _json_safe(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def _decode_json(raw: Any, label: str) -> dict[str, Any]:
    if isinstance(raw, dict):
        value = raw
    else:
        try:
            value = json.loads(str(raw))
        except (TypeError, json.JSONDecodeError) as exc:
            raise ValueError(f"Invalid JSON object for {label}: {exc}") from exc
    if not isinstance(value, dict):
        raise ValueError(f"{label} must be a JSON object")
    return value


def _load_app_config(export_rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], dict[str, int], set[str]]:
    if len(export_rows) != EXPECTED_KEYS:
        raise ValueError(
            f"APP_CONFIG_EXPORT must contain {EXPECTED_KEYS} rows; got {len(export_rows)}"
        )

    payloads: list[dict[str, Any]] = []
    seen: set[str] = set()
    counts = {kind: 0 for kind in EXPECTED_COUNTS}

    for row in export_rows:
        key = _text(_pick(row, "key"))
        kind = _text(_pick(row, "type")).lower()
        status = _text(_pick(row, "sync_status")).upper()

        if not key:
            raise ValueError("APP_CONFIG_EXPORT contains an empty key")
        if key in seen:
            raise ValueError(f"Duplicate app_config key: {key}")
        if kind not in counts:
            raise ValueError(f"Unsupported app_config type for {key}: {kind}")
        if status and status != "READY":
            raise ValueError(f"app_config key {key} is not READY: {status}")

        seen.add(key)
        counts[kind] += 1
        raw_value = _pick(row, "value_jsonb", "value")
        payloads.append({"key": key, "value": _decode_json(raw_value, key)})

    if counts != EXPECTED_COUNTS:
        raise ValueError(f"APP_CONFIG_EXPORT type counts mismatch: {counts}")

    required = {
        "global",
        "template:BASIC",
        "template:PRO",
        "template:ENTERPRISE",
        "tenant:tenant_001",
        "tenant:tenant_001:service:pos",
        "device:infinix_001_410",
    }
    missing = required.difference(seen)
    if missing:
        raise ValueError(f"Missing required app_config keys: {sorted(missing)}")

    _assert_canonical_key_manifest("APP_CONFIG_EXPORT", seen)
    return payloads, counts, seen


def _service_name(row: dict[str, Any]) -> str:
    raw = _text(_pick(row, "service", "service_id", "service_type", "layanan"))
    if not raw:
        key = _text(_pick(row, "key", "app_config_key"))
        marker = ":service:"
        if marker in key:
            raw = key.split(marker, 1)[1]
    if not raw:
        raise ValueError("SERVICES row is missing service/service_id/service_type")
    return raw.strip().lower()


def _build_services(
    rows: list[dict[str, Any]],
    app_config_keys: set[str],
) -> list[dict[str, Any]]:
    if len(rows) != EXPECTED_COUNTS["service"]:
        raise ValueError(
            f"SERVICES must contain exactly {EXPECTED_COUNTS['service']} rows; got {len(rows)}"
        )

    payloads: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()
    for row in rows:
        tenant_id = _required_text(row, "SERVICES.tenant_id", "tenant_id", "tenant", "client_id")
        service = _service_name(row)
        pair = (tenant_id, service)
        if pair in seen:
            raise ValueError(f"Duplicate SERVICES mapping: {tenant_id}/{service}")
        seen.add(pair)

        app_config_key = f"tenant:{tenant_id}:service:{service}"
        explicit_key = _text(_pick(row, "app_config_key", "key"))
        if explicit_key and ":service:" in explicit_key:
            app_config_key = explicit_key
        if app_config_key not in app_config_keys:
            raise ValueError(
                f"SERVICES row {tenant_id}/{service} has no matching app_config key {app_config_key}"
            )

        status = _text(_pick(row, "status", "service_status")) or "active"
        enabled = _bool(
            _pick(row, "enabled", "is_enabled", "service_enabled", "is_active"),
            default=status.lower() not in {"inactive", "disabled", "suspended"},
        )
        payloads.append(
            {
                "tenant_id": tenant_id,
                "service": service,
                "enabled": enabled,
                "status": status.lower(),
                "app_config_key": app_config_key,
            }
        )
    return payloads


def _build_tenants(
    rows: list[dict[str, Any]],
    app_config: list[dict[str, Any]],
    services: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    if len(rows) != EXPECTED_COUNTS["tenant"]:
        raise ValueError(
            f"TENANTS_MASTER must contain exactly {EXPECTED_COUNTS['tenant']} tenants; got {len(rows)}"
        )

    tenant_config = {
        payload["key"].split(":", 1)[1]: payload["value"]
        for payload in app_config
        if payload["key"].startswith("tenant:") and ":service:" not in payload["key"]
    }
    enabled_services: dict[str, set[str]] = {}
    for service in services:
        if service["enabled"]:
            enabled_services.setdefault(service["tenant_id"], set()).add(service["service"])

    payloads: list[dict[str, Any]] = []
    seen: set[str] = set()
    for row in rows:
        tenant_id = _required_text(
            row,
            "TENANTS_MASTER.tenant_id",
            "tenant_id",
            "id",
            "tenant",
            "client_id",
        )
        if tenant_id in seen:
            raise ValueError(f"Duplicate TENANTS_MASTER tenant_id: {tenant_id}")
        seen.add(tenant_id)

        cfg = tenant_config.get(tenant_id, {})
        name = (
            _text(_pick(row, "toko_name", "tenant_name", "name", "store_name"))
            or _text(cfg.get("toko_name"))
            or _text(cfg.get("name"))
            or tenant_id
        )
        template_id = (
            _text(_pick(row, "template_id", "template", "paket", "plan"))
            or _text(cfg.get("template_id"))
            or _text(cfg.get("template"))
            or _text(cfg.get("paket"))
        ).upper()
        if template_id not in EXPECTED_TEMPLATES:
            raise ValueError(
                f"TENANTS_MASTER {tenant_id} has unsupported template_id {template_id!r}"
            )

        raw_status = _text(_pick(row, "status", "tenant_status"))
        if not raw_status:
            raw_status = "suspended" if _bool(_pick(row, "is_blocked"), default=False) else "active"
        status = raw_status.lower()

        configured_features = cfg.get("features")
        if isinstance(configured_features, dict):
            legacy_features = dict(configured_features)
        else:
            services_for_tenant = enabled_services.get(tenant_id, set())
            legacy_features = {
                "pos": "pos" in services_for_tenant,
                "akuntansi": "akuntansi" in services_for_tenant,
            }

        payload: dict[str, Any] = {
            "tenant_id": tenant_id,
            "toko_name": name,
            "template_id": template_id,
            "paket": template_id.lower(),
            # Deprecated compatibility projection only. Runtime reads app_config.
            "features": legacy_features,
            "status": status,
            "is_blocked": status in {"blocked", "suspended", "inactive"},
        }

        optional_columns = {
            "alamat": ("alamat", "address", "toko_address"),
            "package_id": ("package_id",),
            "max_user": ("max_user", "max_users"),
            "expiry": ("expiry", "expiry_date"),
        }
        for target, aliases in optional_columns.items():
            value = _pick(row, *aliases)
            if value is not None and _text(value) != "":
                payload[target] = _json_safe(value)

        payloads.append(payload)

    if set(tenant_config) - seen:
        extra = sorted(set(tenant_config) - seen)[:10]
        raise ValueError(f"app_config contains tenants absent from TENANTS_MASTER: {extra}")

    return payloads


def _build_devices(
    rows: list[dict[str, Any]],
    tenant_ids: set[str],
    app_config_keys: set[str],
) -> list[dict[str, Any]]:
    if len(rows) != EXPECTED_COUNTS["device"]:
        raise ValueError(
            f"DEVICES must contain exactly {EXPECTED_COUNTS['device']} devices; got {len(rows)}"
        )

    payloads: list[dict[str, Any]] = []
    seen: set[str] = set()
    for row in rows:
        device_id = _required_text(row, "DEVICES.device_id", "device_id", "id", "device")
        tenant_id = _required_text(row, "DEVICES.tenant_id", "tenant_id", "tenant", "client_id")
        if device_id in seen:
            raise ValueError(f"Duplicate DEVICES device_id: {device_id}")
        if tenant_id not in tenant_ids:
            raise ValueError(f"DEVICES {device_id} references unknown tenant {tenant_id}")
        seen.add(device_id)

        app_config_key = f"device:{device_id}"
        if app_config_key not in app_config_keys:
            raise ValueError(f"DEVICES {device_id} has no matching app_config key {app_config_key}")

        payload: dict[str, Any] = {
            "device_id": device_id,
            "tenant_id": tenant_id,
            "platform": _text(_pick(row, "platform")) or "android_direct",
            "package_id": _text(_pick(row, "package_id")) or "com.btp.uat",
            "hardware_mode": (_text(_pick(row, "hardware_mode")) or "SIMULATED").upper(),
        }
        optional_columns = {
            "toko_name": ("toko_name", "device_name", "name"),
            "printer_type": ("printer_type",),
            "drawer_mode": ("drawer_mode",),
            "app_version": ("app_version",),
        }
        for target, aliases in optional_columns.items():
            value = _pick(row, *aliases)
            if value is not None and _text(value) != "":
                payload[target] = _json_safe(value)
        payloads.append(payload)

    return payloads


def load_workbook_payloads() -> WorkbookPayloads:
    if not EXCEL_PATH.exists():
        raise FileNotFoundError(f"Excel not found: {EXCEL_PATH}")

    wb = load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    actual_sheets = list(wb.sheetnames)
    missing = [sheet for sheet in EXPECTED_SHEETS if sheet not in actual_sheets]
    unexpected = [sheet for sheet in actual_sheets if sheet not in EXPECTED_SHEETS]
    if missing or unexpected or len(actual_sheets) != len(EXPECTED_SHEETS):
        raise ValueError(
            "Unexpected sheet layout: "
            f"actual={actual_sheets}, missing={missing}, unexpected={unexpected}"
        )

    template_rows = _rows_as_dicts(wb["TEMPLATES"])
    template_ids = {
        _text(_pick(row, "template_id", "template")).upper()
        for row in template_rows
    }
    if template_ids != EXPECTED_TEMPLATES or len(template_rows) != EXPECTED_COUNTS["template"]:
        raise ValueError(f"Unexpected template IDs: {sorted(template_ids)}")

    global_rows = _rows_as_dicts(wb["GLOBAL"])
    if not global_rows:
        raise ValueError("GLOBAL sheet is empty")

    tenant_rows = _rows_as_dicts(wb["TENANTS_MASTER"])
    device_rows = _rows_as_dicts(wb["DEVICES"])
    service_rows = _rows_as_dicts(wb["SERVICES"])
    export_rows = _rows_as_dicts(wb["APP_CONFIG_EXPORT"])

    app_config, counts, app_config_keys = _load_app_config(export_rows)
    services = _build_services(service_rows, app_config_keys)
    tenants = _build_tenants(tenant_rows, app_config, services)
    tenant_ids = {payload["tenant_id"] for payload in tenants}
    devices = _build_devices(device_rows, tenant_ids, app_config_keys)

    if len(tenants) != EXPECTED_COUNTS["tenant"]:
        raise ValueError("Relational tenant payload count mismatch")
    if len(services) != EXPECTED_COUNTS["service"]:
        raise ValueError("Relational service payload count mismatch")
    if len(devices) != EXPECTED_COUNTS["device"]:
        raise ValueError("Relational device payload count mismatch")

    return WorkbookPayloads(
        app_config=app_config,
        tenants=tenants,
        devices=devices,
        services=services,
        counts=counts,
    )


def _headers(service_key: str) -> dict[str, str]:
    return {
        "apikey": service_key,
        "Authorization": f"Bearer {service_key}",
        "Content-Type": "application/json",
        "Accept": "application/json",
        "Prefer": "resolution=merge-duplicates,return=minimal",
        "User-Agent": "berkahtigaputeri-final-pos-sync/3.0",
    }


def _upsert(
    *,
    table: str,
    rows: list[dict[str, Any]],
    on_conflict: str,
    url: str,
    service_key: str,
) -> None:
    endpoint = f"{url.rstrip('/')}/rest/v1/{table}"
    headers = _headers(service_key)
    for start in range(0, len(rows), BATCH_SIZE):
        batch = rows[start : start + BATCH_SIZE]
        response = requests.post(
            endpoint,
            params={"on_conflict": on_conflict},
            headers=headers,
            json=batch,
            timeout=REQUEST_TIMEOUT_SECONDS,
        )
        if not response.ok:
            detail = response.text.strip().replace("\n", " ")[:1000]
            raise RuntimeError(
                f"Supabase REST upsert failed for {table}: "
                f"HTTP {response.status_code}: {detail}"
            )


def _fetch_rows(
    *,
    table: str,
    select: str,
    url: str,
    service_key: str,
) -> list[dict[str, Any]]:
    response = requests.get(
        f"{url.rstrip('/')}/rest/v1/{table}",
        params={"select": select, "limit": "1000"},
        headers={
            "apikey": service_key,
            "Authorization": f"Bearer {service_key}",
            "Accept": "application/json",
            "User-Agent": "berkahtigaputeri-final-pos-sync/3.0",
        },
        timeout=REQUEST_TIMEOUT_SECONDS,
    )
    if not response.ok:
        detail = response.text.strip().replace("\n", " ")[:1000]
        raise RuntimeError(
            f"Supabase REST verification failed for {table}: "
            f"HTTP {response.status_code}: {detail}"
        )
    data = response.json()
    if not isinstance(data, list):
        raise RuntimeError(f"Unexpected verification response for {table}")
    return [dict(item) for item in data]


def _assert_exact_inventory(label: str, expected: Iterable[Any], actual: Iterable[Any]) -> None:
    expected_set = set(expected)
    actual_set = set(actual)
    missing = expected_set - actual_set
    extra = actual_set - expected_set
    if missing or extra:
        missing_sample = sorted(str(value) for value in missing)[:10]
        extra_sample = sorted(str(value) for value in extra)[:10]
        raise RuntimeError(
            f"Post-sync verification failed for {label}; "
            f"missing={missing_sample}; extra={extra_sample}"
        )


def verify_remote(payloads: WorkbookPayloads, url: str, service_key: str) -> None:
    app_rows = _fetch_rows(table="app_config", select="key", url=url, service_key=service_key)
    tenant_rows = _fetch_rows(table="tenants", select="tenant_id", url=url, service_key=service_key)
    device_rows = _fetch_rows(table="devices", select="device_id", url=url, service_key=service_key)
    service_rows = _fetch_rows(
        table="tenant_services",
        select="tenant_id,service",
        url=url,
        service_key=service_key,
    )

    expected_app_keys = [str(row["key"]) for row in payloads.app_config]
    actual_app_keys = [str(row["key"]) for row in app_rows if row.get("key") is not None]
    _assert_exact_inventory("app_config", expected_app_keys, actual_app_keys)
    digest = _assert_canonical_key_manifest("Post-sync remote app_config", actual_app_keys)
    print(f"VERIFIED canonical app_config key manifest MD5: {digest}")
    _assert_exact_inventory(
        "tenants",
        (row["tenant_id"] for row in payloads.tenants),
        (row.get("tenant_id") for row in tenant_rows),
    )
    _assert_exact_inventory(
        "devices",
        (row["device_id"] for row in payloads.devices),
        (row.get("device_id") for row in device_rows),
    )
    _assert_exact_inventory(
        "tenant_services",
        ((row["tenant_id"], row["service"]) for row in payloads.services),
        ((row.get("tenant_id"), row.get("service")) for row in service_rows),
    )

    print(
        "VERIFIED exact workbook-owned inventories: "
        f"app_config={len(payloads.app_config)} "
        f"tenants={len(payloads.tenants)} "
        f"services={len(payloads.services)} "
        f"devices={len(payloads.devices)}"
    )


def sync(payloads: WorkbookPayloads, url: str, service_key: str) -> None:
    # Workbook ownership order respects relational foreign keys first.
    _upsert(
        table="tenants",
        rows=payloads.tenants,
        on_conflict="tenant_id",
        url=url,
        service_key=service_key,
    )
    _upsert(
        table="tenant_services",
        rows=payloads.services,
        on_conflict="tenant_id,service",
        url=url,
        service_key=service_key,
    )
    _upsert(
        table="devices",
        rows=payloads.devices,
        on_conflict="device_id",
        url=url,
        service_key=service_key,
    )
    _upsert(
        table="app_config",
        rows=payloads.app_config,
        on_conflict="key",
        url=url,
        service_key=service_key,
    )
    verify_remote(payloads, url, service_key)


def _print_inventory(payloads: WorkbookPayloads) -> None:
    print(f"Workbook: {EXCEL_PATH}")
    print(f"Sheets validated ({len(EXPECTED_SHEETS)}): {', '.join(EXPECTED_SHEETS)}")
    print(
        "Inventory: "
        f"{payloads.counts['global']} global + "
        f"{payloads.counts['template']} template + "
        f"{payloads.counts['tenant']} tenant + "
        f"{payloads.counts['service']} service + "
        f"{payloads.counts['device']} device = "
        f"{len(payloads.app_config)}"
    )
    print(
        "Relational payloads: "
        f"tenants={len(payloads.tenants)} "
        f"services={len(payloads.services)} "
        f"devices={len(payloads.devices)}"
    )
    print("352 keys validated")
    digest = _assert_canonical_key_manifest("Workbook", (row["key"] for row in payloads.app_config))
    print(f"Canonical key manifest MD5 verified: {digest}")


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Validate the workbook and relational payloads without Supabase credentials or writes.",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)

    try:
        payloads = load_workbook_payloads()
        _print_inventory(payloads)
    except (FileNotFoundError, ValueError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1

    if args.dry_run:
        print("DRY-RUN OK: no Supabase credentials required; no network writes performed")
        return 0

    url = os.environ.get("SUPABASE_URL", "").strip()
    service_key = os.environ.get("SUPABASE_SERVICE_KEY", "").strip()
    if not url or not service_key:
        print(
            "ERROR: SUPABASE_URL and SUPABASE_SERVICE_KEY are required for full sync. "
            "Use --dry-run for validation-only mode.",
            file=sys.stderr,
        )
        return 2

    try:
        sync(payloads, url, service_key)
    except (RuntimeError, requests.RequestException, ValueError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1

    print(
        "DONE Synced workbook ownership: "
        f"{len(payloads.app_config)} app_config keys, "
        f"{len(payloads.tenants)} tenants, "
        f"{len(payloads.services)} services, "
        f"{len(payloads.devices)} devices"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
