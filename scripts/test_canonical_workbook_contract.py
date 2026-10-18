#!/usr/bin/env python3
"""Negative and determinism tests for export_workbook_for_laravel.py.

Fixtures contain synthetic IDs and values only. These tests never contact a
database or network service and never import the prohibited legacy sync script.
"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

import export_workbook_for_laravel as exporter


def _config(label: str) -> str:
    # Intentionally non-alphabetic property order: the exporter must normalize it.
    return json.dumps({"z_label": label, "version": 1, "a_flag": True}, ensure_ascii=False)


def create_fixture(path: Path) -> Path:
    wb = Workbook()
    readme = wb.active
    readme.title = "README_PRODUCTION"
    readme.append(["Synthetic canonical workbook fixture"])
    readme.append(["Used by isolated unit tests only"])

    templates = wb.create_sheet("TEMPLATES")
    templates.append(["template_id", "description"])
    for template_id in ("BASIC", "PRO", "ENTERPRISE"):
        templates.append([template_id, f"Synthetic {template_id} template"])

    global_sheet = wb.create_sheet("GLOBAL")
    global_sheet.append(["setting", "value"])
    global_sheet.append(["timezone", "Asia/Jakarta"])

    tenant_rows = []
    tenant_templates = {}
    for index in range(1, 101):
        tenant_id = f"tenant_{index:03d}"
        template_id = ("BASIC", "PRO", "ENTERPRISE")[(index - 1) % 3]
        tenant_templates[tenant_id] = template_id
        tenant_rows.append([tenant_id, template_id])
    tenants = wb.create_sheet("TENANTS_MASTER")
    tenants.append(["tenant_id", "template_id"])
    for row in tenant_rows:
        tenants.append(row)

    service_rows = []
    for index in range(1, 101):
        tenant_id = f"tenant_{index:03d}"
        service_rows.append([tenant_id, "pos", True, f"tenant:{tenant_id}:service:pos"])
        if index <= 40:
            service_rows.append([tenant_id, "akuntansi", True, f"tenant:{tenant_id}:service:akuntansi"])
    services = wb.create_sheet("SERVICES")
    services.append(["tenant_id", "service", "enabled", "app_config_key"])
    for row in service_rows:
        services.append(row)

    device_rows = []
    for index in range(1, 109):
        device_id = f"device_{index:03d}"
        tenant_id = f"tenant_{((index - 1) % 100) + 1:03d}"
        device_rows.append([device_id, tenant_id, True, f"device:{device_id}"])
    devices = wb.create_sheet("DEVICES")
    devices.append(["device_id", "tenant_id", "is_active", "app_config_key"])
    for row in device_rows:
        devices.append(row)

    export_rows = [["global", _config("global"), "READY"]]
    for template_id in ("BASIC", "PRO", "ENTERPRISE"):
        export_rows.append([f"template:{template_id}", _config(template_id), "READY"])
    for tenant_id in tenant_templates:
        export_rows.append([f"tenant:{tenant_id}", _config(tenant_id), "READY"])
    for tenant_id, service, enabled, key in service_rows:
        export_rows.append([key, _config(key), "READY"])
    for device_id, tenant_id, active, key in device_rows:
        export_rows.append([key, _config(key), "READY"])

    exported = wb.create_sheet("APP_CONFIG_EXPORT")
    exported.append(["key", "value_jsonb", "sync_status"])
    for row in export_rows:
        exported.append(row)

    if len(export_rows) != 352:
        raise AssertionError(f"test fixture bug: expected 352 rows, got {len(export_rows)}")
    wb.save(path)
    wb.close()
    return path


def reverse_data_rows(path: Path, sheet_name: str) -> None:
    wb = load_workbook(path)
    ws = wb[sheet_name]
    headers = [cell.value for cell in ws[1]]
    rows = [
        [ws.cell(row=row_number, column=col).value for col in range(1, len(headers) + 1)]
        for row_number in range(2, ws.max_row + 1)
    ]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for row in reversed(rows):
        ws.append(row)
    wb.save(path)
    wb.close()


class CanonicalWorkbookContractTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tempdir = tempfile.TemporaryDirectory(prefix="canonical-workbook-contract-")
        self.addCleanup(self.tempdir.cleanup)
        self.path = create_fixture(Path(self.tempdir.name) / "fixture.xlsx")

    def _manifest(self) -> dict:
        return exporter.build_manifest(self.path, expected_keyset_md5=None)

    def _mutate(self, sheet_name: str, cell: str, value) -> None:
        wb = load_workbook(self.path)
        wb[sheet_name][cell] = value
        wb.save(self.path)
        wb.close()

    def test_fixture_has_exact_352_inventory_and_layer_categories(self) -> None:
        manifest = self._manifest()
        self.assertEqual(manifest["counts"], {
            "global": 1, "template": 3, "tenant": 100, "service": 140, "device": 108
        })
        self.assertEqual(sum(manifest["counts"].values()), 352)
        self.assertEqual(len(manifest["tenants"]), 100)
        self.assertEqual(len(manifest["services"]), 140)
        self.assertEqual(len(manifest["devices"]), 108)
        self.assertEqual({row["id"] for row in manifest["templates"]}, {"BASIC", "PRO", "ENTERPRISE"})

    def test_manifest_keyset_and_serialization_hashes_are_stable(self) -> None:
        first = self._manifest()
        second = self._manifest()
        self.assertEqual(first, second)
        self.assertEqual(exporter.manifest_bytes(first), exporter.manifest_bytes(second))
        self.assertEqual(len(first["keyset_sha256"]), 64)
        self.assertEqual(len(first["source_workbook_sha256"]), 64)
        self.assertEqual(
            first["manifest_sha256"],
            hashlib.sha256(exporter.canonical_json_bytes({
                key: value for key, value in first.items() if key != "manifest_sha256"
            })).hexdigest(),
        )
        self.assertEqual(
            first["keyset_md5"],
            hashlib.md5("\n".join(sorted(
                ["global"]
                + [f"template:{name}" for name in ("BASIC", "PRO", "ENTERPRISE")]
                + [f"tenant:tenant_{index:03d}" for index in range(1, 101)]
                + [
                    f"tenant:tenant_{index:03d}:service:{service}"
                    for index in range(1, 101)
                    for service in (("pos", "akuntansi") if index <= 40 else ("pos",))
                ]
                + [f"device:device_{index:03d}" for index in range(1, 109)]
            )).encode("utf-8")).hexdigest(),
        )

    def test_projected_lists_are_sorted_independently_of_source_row_order(self) -> None:
        first = self._manifest()
        for sheet_name in ("TEMPLATES", "TENANTS_MASTER", "DEVICES", "SERVICES", "APP_CONFIG_EXPORT"):
            reverse_data_rows(self.path, sheet_name)
        second = self._manifest()
        for field in ("templates", "tenants", "services", "devices", "global_config", "counts", "keyset_md5", "keyset_sha256"):
            self.assertEqual(first[field], second[field], field)
        # The source-byte digest and manifest digest should differ because the workbook bytes differ.
        self.assertNotEqual(first["source_workbook_sha256"], second["source_workbook_sha256"])

    def test_wrong_sheet_order_is_rejected(self) -> None:
        wb = load_workbook(self.path)
        first_name = wb.sheetnames[0]
        second_name = wb.sheetnames[1]
        wb._sheets[0], wb._sheets[1] = wb._sheets[1], wb._sheets[0]
        wb.save(self.path)
        wb.close()
        with self.assertRaisesRegex(ValueError, "sheet order/layout"):
            self._manifest()

    def test_duplicate_export_key_is_rejected(self) -> None:
        wb = load_workbook(self.path)
        ws = wb["APP_CONFIG_EXPORT"]
        ws.append([ws.cell(row=2, column=1).value, _config("duplicate"), "READY"])
        wb.save(self.path)
        wb.close()
        with self.assertRaisesRegex(ValueError, "duplicate key"):
            self._manifest()

    def test_malformed_key_grammar_is_rejected(self) -> None:
        self._mutate("APP_CONFIG_EXPORT", "A2", "tenant:tenant_001:service:pos:unexpected")
        with self.assertRaisesRegex(ValueError, "key grammar"):
            self._manifest()

    def test_missing_template_parent_is_rejected(self) -> None:
        self._mutate("TENANTS_MASTER", "B2", "DOES_NOT_EXIST")
        with self.assertRaisesRegex(ValueError, "invalid template parent"):
            self._manifest()

    def test_tenant_parent_reference_is_rejected(self) -> None:
        self._mutate("SERVICES", "A2", "tenant_9999")
        with self.assertRaisesRegex(ValueError, "unknown tenant parent"):
            self._manifest()

    def test_cross_scope_service_key_is_rejected(self) -> None:
        self._mutate("SERVICES", "D2", "tenant:tenant_001:service:akuntansi")
        with self.assertRaisesRegex(ValueError, "crosses tenant/service scope"):
            self._manifest()

    def test_cross_scope_device_key_is_rejected(self) -> None:
        self._mutate("DEVICES", "D2", "device:device_002")
        with self.assertRaisesRegex(ValueError, "crosses device scope"):
            self._manifest()

    def test_blank_and_numeric_identifiers_are_rejected(self) -> None:
        self._mutate("TENANTS_MASTER", "A2", "   ")
        with self.assertRaisesRegex(ValueError, "TENANTS_MASTER.tenant_id"):
            self._manifest()
        self.path = create_fixture(Path(self.tempdir.name) / "fresh-fixture.xlsx")
        self._mutate("TENANTS_MASTER", "A2", 1)
        with self.assertRaisesRegex(ValueError, "must be stored as text"):
            self._manifest()

    def test_duplicate_headers_are_rejected(self) -> None:
        self._mutate("APP_CONFIG_EXPORT", "B1", "key")
        with self.assertRaisesRegex(ValueError, "duplicate header"):
            self._manifest()

    def test_formula_cells_are_rejected_in_configuration_sheets(self) -> None:
        self._mutate("APP_CONFIG_EXPORT", "A2", '="global"')
        with self.assertRaisesRegex(ValueError, "contains a formula"):
            self._manifest()

    def test_duplicate_json_properties_are_rejected(self) -> None:
        self._mutate("APP_CONFIG_EXPORT", "B2", '{"version":1,"version":2}')
        with self.assertRaisesRegex(ValueError, "not valid strict JSON"):
            self._manifest()

    def test_invalid_service_and_boolean_values_are_rejected(self) -> None:
        self._mutate("SERVICES", "B2", "cashless")
        with self.assertRaisesRegex(ValueError, "unsupported service"):
            self._manifest()
        self.path = create_fixture(Path(self.tempdir.name) / "fresh-fixture.xlsx")
        self._mutate("SERVICES", "C2", "sometimes")
        with self.assertRaisesRegex(ValueError, "explicit boolean"):
            self._manifest()

    def test_unlinked_or_orphaned_export_key_is_rejected(self) -> None:
        wb = load_workbook(self.path)
        ws = wb["APP_CONFIG_EXPORT"]
        for row_number in range(2, ws.max_row + 1):
            if ws.cell(row=row_number, column=1).value == "device:device_001":
                ws.cell(row=row_number, column=1).value = "device:orphan_device"
                break
        else:
            self.fail("fixture did not contain the expected device config key")
        wb.save(self.path)
        wb.close()
        with self.assertRaisesRegex(ValueError, "missing canonical device config key"):
            self._manifest()

    def test_validate_only_does_not_write_manifest(self) -> None:
        output_path = Path(self.tempdir.name) / "must-not-exist.json"
        manifest = self._manifest()
        with patch.object(exporter, "build_manifest", return_value=manifest):
            status = exporter.main([
                "--workbook", str(self.path),
                "--output", str(output_path),
                "--validate-only",
            ])
        self.assertEqual(status, 0)
        self.assertFalse(output_path.exists())


if __name__ == "__main__":
    unittest.main(verbosity=2)
