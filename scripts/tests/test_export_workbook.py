"""Offline-only regression tests for the canonical workbook exporter."""

from __future__ import annotations

import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook

from scripts.export_workbook_for_laravel import (
    DEFAULT_WORKBOOK,
    EXPECTED_COUNTS,
    EXPECTED_SHEETS,
    EXPECTED_TOTAL,
    build_manifest,
    canonical_manifest_sha256,
)


class OfflineWorkbookExporterTests(unittest.TestCase):
    def test_canonical_workbook_has_exact_352_entry_inventory(self) -> None:
        manifest = build_manifest(DEFAULT_WORKBOOK)

        self.assertEqual(EXPECTED_COUNTS, manifest["counts"])
        self.assertEqual(EXPECTED_TOTAL, sum(manifest["counts"].values()))
        self.assertEqual(EXPECTED_TOTAL, len(manifest["tenants"]) + len(manifest["templates"]) + len(manifest["services"]) + len(manifest["devices"]) + 1)
        self.assertEqual(manifest["manifest_sha256"], canonical_manifest_sha256(manifest))
        self.assertEqual(64, len(manifest["manifest_sha256"]))

    def test_two_clean_exports_have_identical_manifest_bytes_and_sha256(self) -> None:
        first = build_manifest(DEFAULT_WORKBOOK)
        second = build_manifest(DEFAULT_WORKBOOK)

        first_bytes = (json.dumps(first, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
        second_bytes = (json.dumps(second, ensure_ascii=False, indent=2) + "\n").encode("utf-8")

        self.assertEqual(first_bytes, second_bytes)
        self.assertEqual(first["manifest_sha256"], second["manifest_sha256"])

    def test_rejects_unexpected_sheet_order(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "wrong-order.xlsx"
            workbook = Workbook()
            workbook.active.title = "WRONG_FIRST_SHEET"
            for name in EXPECTED_SHEETS[1:]:
                workbook.create_sheet(name)
            workbook.save(path)
            workbook.close()

            with self.assertRaisesRegex(ValueError, "Unexpected sheet order/layout"):
                build_manifest(path)

    def test_rejects_incomplete_inventory_before_manifest_export(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "incomplete.xlsx"
            workbook = Workbook()
            workbook.active.title = EXPECTED_SHEETS[0]
            for name in EXPECTED_SHEETS[1:]:
                workbook.create_sheet(name)
            workbook["APP_CONFIG_EXPORT"].append(["key", "value_jsonb", "sync_status"])
            workbook.save(path)
            workbook.close()

            with self.assertRaisesRegex(ValueError, "Expected .* keys"):
                build_manifest(path)


if __name__ == "__main__":
    unittest.main()
