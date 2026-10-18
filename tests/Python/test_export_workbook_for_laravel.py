from __future__ import annotations

import ast
import hashlib
import json
import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from scripts.export_workbook_for_laravel import (
    EXPECTED_COUNTS,
    EXPECTED_KEYSET_MD5,
    EXPECTED_SHEETS,
    build_manifest,
    canonical_manifest_bytes,
    normalize_header,
    parse_value,
    write_manifest,
)

ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = ROOT / "config/FINAL_Production_Excel_A_WebC_Template_100Tenants.xlsx"


class CanonicalWorkbookExporterContractTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.fixture = Path(self.temp.name) / WORKBOOK.name
        shutil.copyfile(WORKBOOK, self.fixture)

    def mutate(self, sheet_name: str, mutate_row) -> None:
        workbook = load_workbook(self.fixture)
        try:
            sheet = workbook[sheet_name]
            headers = [normalize_header(cell.value) for cell in sheet[1]]
            mutate_row(sheet, {name: index + 1 for index, name in enumerate(headers) if name})
            workbook.save(self.fixture)
        finally:
            workbook.close()

    def test_canonical_workbook_has_exact_shape_inventory_and_key_hash(self) -> None:
        manifest = build_manifest(WORKBOOK)
        workbook = load_workbook(WORKBOOK, read_only=True)
        try:
            self.assertEqual(EXPECTED_SHEETS, workbook.sheetnames)
        finally:
            workbook.close()
        self.assertEqual(EXPECTED_COUNTS, manifest["counts"])
        self.assertEqual(352, sum(manifest["counts"].values()))
        self.assertEqual(EXPECTED_KEYSET_MD5, manifest["keyset_md5"])
        self.assertEqual("2df1a708480e039fef0111f49da015beeb18de7d3896e16fd3debf74cdd602cd", manifest["source_sha256"])
        self.assertEqual(hashlib.sha256(WORKBOOK.read_bytes()).hexdigest(), manifest["source_sha256"])
        self.assertEqual(3, len(manifest["templates"]))
        self.assertEqual(100, len(manifest["tenants"]))
        self.assertEqual(140, len(manifest["services"]))
        self.assertEqual(108, len(manifest["devices"]))

    def test_duplicate_export_key_fails_closed(self) -> None:
        def duplicate(sheet, columns):
            key_col = columns["key"]
            sheet.cell(row=3, column=key_col).value = sheet.cell(row=2, column=key_col).value

        self.mutate("APP_CONFIG_EXPORT", duplicate)
        with self.assertRaisesRegex(ValueError, "Duplicate canonical config key"):
            build_manifest(self.fixture)

    def test_missing_or_replaced_export_key_fails_closed(self) -> None:
        def replace_key(sheet, columns):
            sheet.cell(row=sheet.max_row, column=columns["key"]).value = "tenant:tenant_missing"

        self.mutate("APP_CONFIG_EXPORT", replace_key)
        with self.assertRaises(ValueError):
            build_manifest(self.fixture)

    def test_non_normalized_tenant_identifier_fails_closed(self) -> None:
        def uppercase_first_tenant(sheet, columns):
            row = 2
            original = sheet.cell(row=row, column=columns["tenant_id"]).value
            sheet.cell(row=row, column=columns["tenant_id"]).value = str(original).upper()

        self.mutate("TENANTS_MASTER", uppercase_first_tenant)
        with self.assertRaisesRegex(ValueError, "identifier is not in canonical normalized form"):
            build_manifest(self.fixture)

    def test_duplicate_tenant_identifier_fails_closed(self) -> None:
        def duplicate_tenant(sheet, columns):
            sheet.cell(row=3, column=columns["tenant_id"]).value = sheet.cell(row=2, column=columns["tenant_id"]).value

        self.mutate("TENANTS_MASTER", duplicate_tenant)
        with self.assertRaisesRegex(ValueError, "Empty/duplicate tenant_id"):
            build_manifest(self.fixture)

    def test_service_mapping_must_match_its_exact_key(self) -> None:
        def change_service_key(sheet, columns):
            sheet.cell(row=2, column=columns["app_config_key"]).value = "global"

        self.mutate("SERVICES", change_service_key)
        with self.assertRaisesRegex(ValueError, "exact canonical key"):
            build_manifest(self.fixture)

    def test_relational_scope_totals_must_match_export(self) -> None:
        def remove_last_service(sheet, columns):
            rows = [
                row for row in range(2, sheet.max_row + 1)
                if any(sheet.cell(row=row, column=column).value is not None for column in range(1, sheet.max_column + 1))
            ]
            row = rows[-1]
            for column in range(1, sheet.max_column + 1):
                sheet.cell(row=row, column=column).value = None

        self.mutate("SERVICES", remove_last_service)
        with self.assertRaisesRegex(ValueError, "SERVICES must contain 140 rows"):
            build_manifest(self.fixture)

    def test_unexpected_sheet_layout_fails_closed(self) -> None:
        workbook = load_workbook(self.fixture)
        try:
            workbook["GLOBAL"].title = "GLOBAL_OLD"
            workbook.save(self.fixture)
        finally:
            workbook.close()
        with self.assertRaisesRegex(ValueError, "Unexpected sheet order/layout"):
            build_manifest(self.fixture)

    def test_export_scope_type_must_match_key_prefix(self) -> None:
        def wrong_type(sheet, columns):
            sheet.cell(row=2, column=columns["type"]).value = "device"

        self.mutate("APP_CONFIG_EXPORT", wrong_type)
        with self.assertRaisesRegex(ValueError, "type=.*does not match expected"):
            build_manifest(self.fixture)

    def test_required_boolean_flag_rejects_empty_and_unknown_values(self) -> None:
        from scripts.export_workbook_for_laravel import bool_value

        for malformed in (None, "", "sometimes", 2):
            with self.subTest(malformed=malformed):
                with self.assertRaises(ValueError):
                    bool_value(malformed)

    def test_malformed_json_fails_closed(self) -> None:
        def malformed(sheet, columns):
            sheet.cell(row=2, column=columns["value_jsonb"]).value = "{ malformed"

        self.mutate("APP_CONFIG_EXPORT", malformed)
        with self.assertRaisesRegex(ValueError, "not valid JSON"):
            build_manifest(self.fixture)

    def test_explicit_null_false_zero_and_empty_values_are_preserved(self) -> None:
        parsed = parse_value(
            json.dumps({"null_value": None, "false_value": False, "zero_value": 0, "empty_value": ""}),
            "fixture",
        )
        self.assertIsNone(parsed["null_value"])
        self.assertIs(parsed["false_value"], False)
        self.assertEqual(0, parsed["zero_value"])
        self.assertEqual("", parsed["empty_value"])
        self.assertEqual(1, parsed["version"])


    def test_manifest_projection_order_is_stable(self) -> None:
        manifest = build_manifest(WORKBOOK)
        self.assertEqual(sorted(row["tenant_id"] for row in manifest["tenants"]), [row["tenant_id"] for row in manifest["tenants"]])
        self.assertEqual(sorted((row["tenant_id"], row["service"]) for row in manifest["services"]), [(row["tenant_id"], row["service"]) for row in manifest["services"]])
        self.assertEqual(sorted(row["device_id"] for row in manifest["devices"]), [row["device_id"] for row in manifest["devices"]])

    def test_canonical_manifest_bytes_and_sha256_repeat_across_clean_runs(self) -> None:
        first_manifest = build_manifest(WORKBOOK)
        second_manifest = build_manifest(WORKBOOK)
        with tempfile.TemporaryDirectory() as first_dir, tempfile.TemporaryDirectory() as second_dir:
            first_path = Path(first_dir) / "manifest.json"
            second_path = Path(second_dir) / "manifest.json"
            first_digest = write_manifest(first_manifest, first_path)
            second_digest = write_manifest(second_manifest, second_path)
            first_bytes = first_path.read_bytes()
            second_bytes = second_path.read_bytes()
            self.assertEqual(first_bytes, second_bytes)
            self.assertEqual(canonical_manifest_bytes(first_manifest), first_bytes)
            self.assertEqual(hashlib.sha256(first_bytes).hexdigest(), first_digest)
            self.assertEqual(first_digest, second_digest)
            self.assertEqual(f"{first_digest}  manifest.json\n".encode("ascii"), first_path.with_name("manifest.json.sha256").read_bytes())
            self.assertEqual(first_path.with_name("manifest.json.sha256").read_bytes(), second_path.with_name("manifest.json.sha256").read_bytes())

    def test_offline_exporter_imports_no_network_or_database_clients(self) -> None:
        source_path = ROOT / "scripts/export_workbook_for_laravel.py"
        source = source_path.read_text(encoding="utf-8")
        tree = ast.parse(source, filename=str(source_path))
        imported_roots: set[str] = set()
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                imported_roots.update(alias.name.split(".", 1)[0] for alias in node.names)
            elif isinstance(node, ast.ImportFrom) and node.module:
                imported_roots.add(node.module.split(".", 1)[0])
        forbidden = {"requests", "httpx", "urllib", "urllib3", "psycopg", "psycopg2", "sqlalchemy", "supabase", "socket", "sqlite3", "asyncpg"}
        self.assertFalse(imported_roots & forbidden, sorted(imported_roots & forbidden))
        self.assertNotIn("SUPABASE_URL", source)
        self.assertNotIn("SUPABASE_SERVICE_KEY", source)

    def test_extra_export_row_fails_closed(self) -> None:
        def append_extra(sheet, columns):
            row = sheet.max_row + 1
            sheet.cell(row=row, column=columns["key"]).value = "tenant:tenant_extra"
            sheet.cell(row=row, column=columns["type"]).value = "tenant"
            sheet.cell(row=row, column=columns["value_jsonb"]).value = json.dumps({"version": 1})
            sheet.cell(row=row, column=columns["sync_status"]).value = "READY"
        self.mutate("APP_CONFIG_EXPORT", append_extra)
        with self.assertRaisesRegex(ValueError, "Expected.*keys"):
            build_manifest(self.fixture)

    def test_missing_template_assignment_fails_closed(self) -> None:
        def remove_assignment(sheet, columns):
            sheet.cell(row=2, column=columns["template_id"]).value = ""
        self.mutate("TENANTS_MASTER", remove_assignment)
        with self.assertRaisesRegex(ValueError, "template_id must be a canonical uppercase template ID"):
            build_manifest(self.fixture)

    def test_malformed_device_identifier_fails_closed(self) -> None:
        def malformed_device(sheet, columns):
            sheet.cell(row=2, column=columns["device_id"]).value = "Device Bad"
        self.mutate("DEVICES", malformed_device)
        with self.assertRaisesRegex(ValueError, "identifier is not in canonical normalized form"):
            build_manifest(self.fixture)


if __name__ == "__main__":
    unittest.main()
